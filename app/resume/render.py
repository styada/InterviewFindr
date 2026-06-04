"""Render structured resumes and cover letters to PDF and DOCX bytes.

The orchestrator (app.pipeline.run) calls these after Stage 5 tailoring.
PDFs are produced by WeasyPrint from Jinja2 templates in this package's
``templates/`` directory; DOCX is built with python-docx. All renderers
are pure functions: they return bytes and do not touch the database or
object store — the orchestrator decides where the bytes go.
"""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from jinja2 import Environment, FileSystemLoader
from weasyprint import HTML

log = logging.getLogger(__name__)

_TEMPLATE_DIR = Path(__file__).parent / "templates"
_env: Environment = Environment(
    loader=FileSystemLoader(str(_TEMPLATE_DIR)),
    autoescape=True,
)


def render_resume_pdf(resume: dict[str, Any]) -> bytes:
    """Render a structured resume dict to a PDF byte string."""
    ctx: dict[str, Any] = {
        "contact": resume.get("contact") or {},
        "summary": resume.get("summary", "") or "",
        "experience": resume.get("experience") or [],
        "education": resume.get("education") or [],
        "skills": resume.get("skills") or [],
    }
    html = _env.get_template("resume.html").render(**ctx)
    return HTML(string=html).write_pdf()


def render_resume_docx(resume: dict[str, Any]) -> bytes:
    """Render a structured resume dict to a DOCX byte string."""
    doc = Document()
    contact = resume.get("contact") or {}

    name = str(contact.get("name") or "")
    if name:
        doc.add_heading(name, level=0)

    contact_parts: list[str] = []
    for key in ("email", "phone", "location"):
        v = contact.get(key)
        if v:
            contact_parts.append(str(v))
    if contact_parts:
        p = doc.add_paragraph(" · ".join(contact_parts))
        for run in p.runs:
            run.font.size = Pt(9)

    summary = resume.get("summary")
    if summary:
        p = doc.add_paragraph()
        run = p.add_run(str(summary))
        run.italic = True

    experience = resume.get("experience") or []
    if experience:
        doc.add_heading("Experience", level=1)
        for r in experience:
            if not isinstance(r, dict):
                continue
            title = str(r.get("title") or "")
            company = str(r.get("company") or "")
            p = doc.add_paragraph()
            run = p.add_run(f"{title} — {company}")
            run.bold = True

            meta_parts: list[str] = []
            start = r.get("start") or ""
            end = r.get("end") or ""
            if start or end:
                meta_parts.append(f"{start} – {end}")
            if r.get("location"):
                meta_parts.append(str(r["location"]))
            if meta_parts:
                meta_p = doc.add_paragraph(" · ".join(meta_parts))
                for mr in meta_p.runs:
                    mr.font.size = Pt(9)

            for b in r.get("bullets") or []:
                doc.add_paragraph(str(b), style="List Bullet")

    education = resume.get("education") or []
    if education:
        doc.add_heading("Education", level=1)
        for e in education:
            if not isinstance(e, dict):
                continue
            line = str(e.get("degree") or "")
            if e.get("field"):
                line += f", {e['field']}"
            line += f" — {e.get('school') or ''}"
            if e.get("year"):
                line += f" ({e['year']})"
            doc.add_paragraph(line)

    skills = resume.get("skills") or []
    if skills:
        doc.add_heading("Skills", level=1)
        names: list[str] = []
        for s in skills:
            if isinstance(s, dict) and s.get("name"):
                names.append(str(s["name"]))
            elif isinstance(s, str):
                names.append(s)
        doc.add_paragraph(", ".join(names))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_cover_letter_pdf(text: str) -> bytes:
    """Render a cover letter (paragraphs separated by blank lines) to PDF bytes."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    html = _env.get_template("cover_letter.html").render(paragraphs=paragraphs)
    return HTML(string=html).write_pdf()


def resume_to_plain_text(resume: dict[str, Any]) -> str:
    """Flatten a structured resume to a single string for keyword validation.

    Used by Stage 6 (ATS validation) to check that the tailored resume
    contains the must-have keywords from the parsed job description.
    """
    out: list[str] = []
    summary = resume.get("summary")
    if summary:
        out.append(str(summary))
    for r in resume.get("experience") or []:
        if not isinstance(r, dict):
            continue
        out.append(f"{r.get('title', '')} at {r.get('company', '')}")
        for b in r.get("bullets") or []:
            out.append(str(b))
    for s in resume.get("skills") or []:
        if isinstance(s, dict) and s.get("name"):
            out.append(str(s["name"]))
        elif isinstance(s, str):
            out.append(s)
    return "\n".join(out)
